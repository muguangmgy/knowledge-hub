from __future__ import annotations

import re
import statistics
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

from app.common.errors import AppError, ErrorCode

ALLOWED_TYPES = frozenset({"pdf", "docx", "md", "txt"})

# docx 启发式标题：样式未命中时的兜底闸门
_DOCX_HEADING_MAX_CHARS = 40
_DOCX_HEADING_END_MARKS = "。！？；.!?"
_DOCX_BOLD_RATIO = 0.8
_DOCX_SIZE_LARGE_RATIO = 1.15
_DOCX_SIZE_H1_RATIO = 1.3
_DOCX_DEFAULT_BODY_PT = 12.0
_DOCX_TOC_OR_DECOR = re.compile(r"\.{3,}|\u2026|\t")
_DOCX_PURE_DECOR = re.compile(r"^[\d\s.·•\-—_*]+$")

# pdf 无书签时的字号启发式（相对正文中位，不依赖中文节名关键词）
_PDF_HEADING_MAX_CHARS = 40
_PDF_HEADING_END_MARKS = _DOCX_HEADING_END_MARKS
_PDF_SIZE_SUB_RATIO = 1.15  # 二级标题下限（介于正文与节标题之间）
_PDF_SIZE_SECTION_RATIO = 1.4  # 一级节标题下限
_PDF_SIZE_TITLE_RATIO = 1.7  # 封面大标题/姓名等超大字不当节标题（正文中位约 1.4~1.5 倍才是节名）
_PDF_DEFAULT_BODY_PT = 10.0
# 项目符号、日期开头、短标签字段（报告日期：）不当标题
_PDF_NOT_HEADING_PREFIX = re.compile(
    r"^(?:[•·●○◦▪▫]\s*|https?://|\d{4}[-/.]\d{1,2}[-/.]\d{1,2}\b|[^：:]{1,12}[：:])"
)

def detect_file_type(filename: str) -> str:
    """白名单校验；不支持类型（含 .doc）→ 42201。"""
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix not in ALLOWED_TYPES:
        raise AppError(
            ErrorCode.UNSUPPORTED_TYPE,
            f"不支持的文件类型: .{suffix or '?'}",
            http_status=422,
        )
    return suffix


def parse_bytes(file_type: str, data: bytes) -> str:
    """按类型提取纯文本；无有效正文由上层回滚为 50001。

    docx：Heading 样式优先，否则加粗/字号启发式。
    pdf：书签（toc）优先；无书签则按字号启发式（视觉标题 → #/##）。
    均转为 Markdown # 行，便于 chunker 填 heading_path。
    """
    try:
        if file_type == "pdf":
            return _parse_pdf(data)
        if file_type == "docx":
            return _parse_docx(data)
        if file_type in {"md", "txt"}:
            return _decode_text(data)
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            ErrorCode.INGEST_FAILED,
            f"文档解析失败: {exc}",
            http_status=500,
        ) from exc
    raise AppError(ErrorCode.UNSUPPORTED_TYPE, "不支持的文件类型", http_status=422)


def _decode_text(data: bytes) -> str:
    """优先 UTF-8；失败再试 GBK，避免中文 Windows 文本静默丢字。"""
    for encoding in ("utf-8", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    # 仍失败则替换非法字节，保证入库链路可继续、由上层判空
    return data.decode("utf-8", errors="replace")


def _parse_pdf(data: bytes) -> str:
    """PDF → 文本；有书签用 toc，无书签则按字号启发式注入 Markdown 标题。

    启发式跨页保留当前一级节（如「项目经历」），以便次页项目名仍挂到正确路径。
    扫描件/纯图无字时上层会判无正文。不在此处做 OCR。
    """
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        headings_by_page = _pdf_toc_headings(doc)
        if not headings_by_page:
            return _pdf_text_with_heuristic_headings(doc)
        return _pdf_text_with_toc_headings(doc, headings_by_page)
    finally:
        doc.close()


def _pdf_toc_headings(doc) -> dict[int, list[tuple[int, str]]]:
    """将 get_toc() 转为 0-based 页码 → [(level, title), ...]。无书签或读取失败返回 {}。"""
    try:
        toc = doc.get_toc() or []
    except Exception:
        return {}

    by_page: dict[int, list[tuple[int, str]]] = {}
    page_count = int(doc.page_count or 0)
    if page_count <= 0:
        return {}

    for entry in toc:
        if not entry or len(entry) < 3:
            continue
        try:
            level = int(entry[0])
            page = int(entry[2])
        except (TypeError, ValueError):
            continue
        title = str(entry[1] or "").strip()
        if not title or page < 1:
            continue
        # toc 页码 1-based；越界夹到末页，避免丢标题
        idx = min(page - 1, page_count - 1)
        level = max(1, min(6, level))
        by_page.setdefault(idx, []).append((level, title))
    return by_page


def _pdf_text_with_toc_headings(
    doc,
    headings_by_page: dict[int, list[tuple[int, str]]],
) -> str:
    """有书签：在页内把对应行替换成 Markdown 标题，避免节名再出现一次。"""
    parts: list[str] = []
    for i, page in enumerate(doc):
        text = page.get_text("text") or ""
        piece = _pdf_apply_toc_headings(text, headings_by_page.get(i, []))
        if piece:
            parts.append(piece)
    return "\n\n".join(parts).strip()


def _pdf_apply_toc_headings(
    text: str,
    headings: list[tuple[int, str]],
) -> str:
    """将页内与书签一致的行改成 # 标题；未出现在正文里的书签补在页首。"""
    if not headings:
        return text.strip()

    lines = text.splitlines()
    used: set[int] = set()
    for idx, line in enumerate(lines):
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        for hi, (level, title) in enumerate(headings):
            if hi in used:
                continue
            if not _pdf_toc_title_matches_line(title, stripped):
                continue
            lines[idx] = f"{'#' * level} {title.strip()}"
            used.add(hi)
            break

    missing = [
        f"{'#' * level} {title.strip()}"
        for hi, (level, title) in enumerate(headings)
        if hi not in used
    ]
    body = "\n".join(lines).strip()
    if missing and body:
        return "\n\n".join(missing) + "\n\n" + body
    if missing:
        return "\n\n".join(missing)
    return body


def _pdf_toc_title_matches_line(title: str, line: str) -> bool:
    """书签名与页面行比对：去空白后相等，或一方为另一方前缀且足够长。"""
    t = " ".join((title or "").split())
    s = " ".join((line or "").split())
    if not t or not s:
        return False
    if t == s:
        return True
    shorter, longer = (t, s) if len(t) <= len(s) else (s, t)
    if len(shorter) < 4:
        return False
    return longer.startswith(shorter) and (len(shorter) / len(longer)) >= 0.7


def _pdf_text_with_heuristic_headings(doc) -> str:
    """全文按视觉阅读顺序 + 字号层级重建 Markdown 标题。

    规则通用：一级=明显大于正文；二级=介于正文与一级之间，且已有一级节时挂上去。
    封面超大字、短标签字段不当节标题，这样文首题名/日期等仍落在「正文」。
    按行抽取（同一行多栏不合并），只把「下一行左对齐短尾巴」拼回上一行（含跨页），
    避免研报长标题折行变成假子节，也不把简历里学校/学历/日期粘成空标题。
    不写死「教育/项目」等词；有 toc 的 PDF 不走此路径。
    """
    rows, body_median_pt = _pdf_collect_lines(doc)
    lines_out: list[str] = []
    has_section = False
    for _page, _y0, _x0, text, size in _merge_pdf_wrap_lines(rows):
        level = _pdf_heuristic_heading_level(
            text, size_pt=size, body_median_pt=body_median_pt
        )
        if level == 1:
            has_section = True
            lines_out.append(f"# {text}")
        elif level == 2 and has_section:
            lines_out.append(f"## {text}")
        else:
            lines_out.append(text)
    return "\n".join(lines_out).strip()


def _pdf_collect_lines(
    doc,
) -> tuple[list[tuple[int, float, float, str, float]], float]:
    """一次遍历抽行，并按非空 span 字号取中位作正文基准。"""
    rows: list[tuple[int, float, float, str, float]] = []
    sizes: list[float] = []
    for page_i, page in enumerate(doc):
        page_rows, page_sizes = _pdf_page_sorted_lines(page)
        sizes.extend(page_sizes)
        rows.extend((page_i, y, x, text, size) for y, x, text, size in page_rows)
    median = (
        float(statistics.median(sizes)) if sizes else _PDF_DEFAULT_BODY_PT
    )
    return rows, median


def _pdf_page_sorted_lines(
    page,
) -> tuple[list[tuple[float, float, str, float]], list[float]]:
    """一次 get_text(dict)：按 (y, x) 排行，并收集非空 span 字号。失败则空。"""
    try:
        data = page.get_text("dict")
    except Exception:
        return [], []

    rows: list[tuple[float, float, str, float]] = []
    sizes: list[float] = []
    for block in data.get("blocks") or []:
        if block.get("type", 0) != 0:
            continue
        for line in block.get("lines") or []:
            spans = line.get("spans") or []
            if not spans:
                continue
            for span in spans:
                piece = (span.get("text") or "").strip()
                if not piece:
                    continue
                span_size = float(span.get("size") or 0)
                if span_size > 0:
                    sizes.append(span_size)
            text = "".join(str(s.get("text") or "") for s in spans).strip()
            if not text:
                continue
            bbox = line.get("bbox") or (0, 0, 0, 0)
            size = max(float(s.get("size") or 0) for s in spans)
            rows.append((float(bbox[1]), float(bbox[0]), text, size))
    rows.sort(key=lambda r: (round(r[0], 1), r[1]))
    return rows, sizes


def _join_pdf_wrapped_text(left: str, right: str) -> str:
    """拼纵向折行：中文直接相连；两侧都是 ASCII 单词时补空格。"""
    if not left:
        return right
    if not right:
        return left
    if (
        left[-1].isascii()
        and left[-1].isalnum()
        and right[0].isascii()
        and right[0].isalnum()
    ):
        return f"{left} {right}"
    return left + right


def _pdf_is_vertical_wrap(
    prev: tuple[int, float, float, str, float] | None,
    curr: tuple[int, float, float, str, float],
) -> bool:
    """下一行左对齐、字号接近、明显更短 → 视为上一行折下的尾巴。

    同一行多栏（学校 | 学历 | 日期，y 几乎相同）不算折行。
    跨页时下一页 y 会回到页顶，只要求相邻页、字号/缩进接近。
    """
    if prev is None:
        return False
    p_page, py, px, ptext, psize = prev
    c_page, cy, cx, ctext, csize = curr
    if not ptext or not ctext:
        return False
    if abs(csize - psize) > 0.6:
        return False
    if abs(cx - px) > 28.0:
        return False
    if ptext[-1] in _PDF_HEADING_END_MARKS:
        return False
    if len(ptext) < 16:
        return False
    if len(ctext) >= max(8, int(len(ptext) * 0.6)):
        return False
    if c_page == p_page:
        if cy - py < 2.0:
            return False
        if cy - py > max(22.0, psize * 1.8):
            return False
    elif c_page == p_page + 1:
        pass
    else:
        return False
    return True


def _merge_pdf_wrap_lines(
    rows: list[tuple[int, float, float, str, float]],
) -> list[tuple[int, float, float, str, float]]:
    """把纵向折行尾巴拼进上一行，供标题长度判断；含跨页短尾巴。"""
    merged: list[tuple[int, float, float, str, float]] = []
    for row in rows:
        if merged and _pdf_is_vertical_wrap(merged[-1], row):
            page, y, x, text, size = merged[-1]
            merged[-1] = (
                page,
                y,
                x,
                _join_pdf_wrapped_text(text, row[3]),
                max(size, row[4]),
            )
        else:
            merged.append(row)
    return merged


def _pdf_heuristic_heading_level(
    text: str,
    *,
    size_pt: float,
    body_median_pt: float,
) -> int | None:
    """返回 1=节标题、2=候选子标题、None=正文。"""
    if len(text) > _PDF_HEADING_MAX_CHARS:
        return None
    if text[-1] in _PDF_HEADING_END_MARKS:
        return None
    if _DOCX_PURE_DECOR.fullmatch(text) or _DOCX_TOC_OR_DECOR.search(text):
        return None
    if _PDF_NOT_HEADING_PREFIX.search(text):
        return None
    if size_pt >= body_median_pt * _PDF_SIZE_TITLE_RATIO:
        return None
    if size_pt >= body_median_pt * _PDF_SIZE_SECTION_RATIO:
        return 1
    if size_pt >= body_median_pt * _PDF_SIZE_SUB_RATIO:
        return 2
    return None


def _parse_docx(data: bytes) -> str:
    """docx → 文本；Heading 样式优先，否则启发式标题；表格单元格一并抽出。"""
    document = DocxDocument(BytesIO(data))
    paras: list[tuple[object, str, float | None]] = []
    sizes: list[float] = []
    for p in document.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        size_pt = _docx_effective_font_pt(p)
        if size_pt is not None:
            sizes.append(size_pt)
        paras.append((p, text, size_pt))

    body_median_pt = (
        float(statistics.median(sizes)) if sizes else _DOCX_DEFAULT_BODY_PT
    )

    lines: list[str] = []
    for p, text, size_pt in paras:
        level = _docx_heading_level(p)
        if level is None:
            level = _docx_heuristic_heading_level(
                p, text, size_pt, body_median_pt=body_median_pt
            )
        if level:
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                lines.append("\t".join(cells))

    return "\n\n".join(lines)


def _docx_heading_level(paragraph) -> int | None:
    """从样式名解析 Heading 1..6 / Title；非标题返回 None。"""
    style = paragraph.style
    if style is None or not style.name:
        return None
    name = style.name.strip()
    # 常见：Heading 1 / 标题 1
    for prefix in ("Heading ", "标题 "):
        if name.startswith(prefix):
            rest = name[len(prefix) :].strip()
            if rest.isdigit():
                level = int(rest)
                if 1 <= level <= 6:
                    return level
    if name in {"Title", "标题"}:
        return 1
    return None


def _docx_heuristic_heading_level(
    paragraph,
    text: str,
    size_pt: float | None,
    *,
    body_median_pt: float,
) -> int | None:
    """无 Heading 样式时的兜底：短行 + 加粗或偏大字号 + 无句末标点。"""
    if len(text) > _DOCX_HEADING_MAX_CHARS:
        return None
    if text[-1] in _DOCX_HEADING_END_MARKS:
        return None
    if _DOCX_PURE_DECOR.fullmatch(text) or _DOCX_TOC_OR_DECOR.search(text):
        return None

    mostly_bold = _docx_mostly_bold(paragraph)
    large = size_pt is not None and size_pt >= body_median_pt * _DOCX_SIZE_LARGE_RATIO
    if not mostly_bold and not large:
        return None

    if size_pt is not None and size_pt >= body_median_pt * _DOCX_SIZE_H1_RATIO:
        return 1
    return 2


def _docx_mostly_bold(paragraph) -> bool:
    """段落可见字符中加粗占比达到阈值则视为整段加粗。"""
    runs = list(paragraph.runs or [])
    if not runs:
        return bool(_docx_style_bold(paragraph))

    bold_chars = 0
    total = 0
    style_bold = _docx_style_bold(paragraph)
    for run in runs:
        piece = run.text or ""
        if not piece:
            continue
        total += len(piece)
        flag = run.bold
        if flag is None:
            flag = style_bold
        if flag:
            bold_chars += len(piece)
    if total <= 0:
        return bool(style_bold)
    return (bold_chars / total) >= _DOCX_BOLD_RATIO


def _docx_style_bold(paragraph) -> bool:
    """沿样式链取 bold；未声明视为 False。"""
    style = paragraph.style
    while style is not None:
        font = getattr(style, "font", None)
        if font is not None and font.bold is not None:
            return bool(font.bold)
        style = getattr(style, "base_style", None)
    return False


def _docx_effective_font_pt(paragraph) -> float | None:
    """段落有效字号（pt）：优先 run 最大字号，否则沿样式链回落。"""
    pts: list[float] = []
    for run in paragraph.runs or []:
        size = getattr(run.font, "size", None) if run.font is not None else None
        if size is not None:
            pts.append(float(size.pt))
    if pts:
        return max(pts)

    style = paragraph.style
    while style is not None:
        font = getattr(style, "font", None)
        if font is not None and font.size is not None:
            return float(font.size.pt)
        style = getattr(style, "base_style", None)
    return None
